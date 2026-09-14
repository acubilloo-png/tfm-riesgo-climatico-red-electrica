# ============================================================================
# 20_quitar_etiqueta_clasificacion.py
# ----------------------------------------------------------------------------
# Retira del documento la etiqueta de confidencialidad corporativa heredada de la
# plantilla y el marcado de contenido que la acompana, que es lo que hace que
# Word estampe "Internal Use" en el pie de cada pagina.
#
# ORIGEN:  20260817_TFMAlvaroCubillo_draftRicardo.docx        (no se modifica)
# DESTINO: 20260817_TFMAlvaroCubillo_draftRicardo_limpio.docx
#
# ----------------------------------------------------------------------------
# QUE ES LA MARCA, EXACTAMENTE
# ----------------------------------------------------------------------------
# No es una marca de agua de Word. En los encabezados y pies no hay ni una forma
# ni una imagen: la que hubo se retiro ya en una revision anterior. Lo que queda
# es la etiqueta de sensibilidad de Microsoft Purview y sus propiedades, que son
# la instruccion para volver a dibujarla:
#
#   docMetadata/LabelInfo.xml   etiqueta aplicada, con contentBits="2", que es
#                               precisamente "marcar el contenido en el pie"
#   ClassificationContentMarkingFooterText        "Internal Use"
#   ClassificationContentMarkingFooterFontProps   #0000ff, 12 pt, Calibri
#   ClassificationContentMarkingFooterShapeIds    los 27 identificadores de las
#                                                 formas que Word vuelve a crear
#
# Es decir: borrar la marca del pie no basta, porque la etiqueta la regenera al
# abrir o guardar. Hay que retirar la etiqueta.
#
# ----------------------------------------------------------------------------
# LO QUE ESTE SCRIPT NO TOCA
# ----------------------------------------------------------------------------
#   · El texto, las tablas, las imagenes y los comentarios: se comprueba al final
#     que no ha cambiado ni un caracter.
#   · ContentTypeId, que es un residuo del sitio de SharePoint de origen y no
#     produce ninguna marca visible.
# ============================================================================

import os
import re
import shutil
import zipfile

import docx
from lxml import etree

CARPETA = os.path.dirname(os.path.abspath(__file__))
ORIGEN = os.path.join(CARPETA, "20260817_TFMAlvaroCubillo_draftRicardo.docx")
DESTINO = os.path.join(CARPETA,
                       "20260817_TFMAlvaroCubillo_draftRicardo_limpio.docx")

PARTE_ETIQUETA = "docMetadata/LabelInfo.xml"
PROPIEDADES = ["ClassificationContentMarkingFooterText",
               "ClassificationContentMarkingFooterFontProps",
               "ClassificationContentMarkingFooterShapeIds",
               "ClassificationContentMarkingFooterShapeIds-1",
               "ClassificationContentMarkingHeaderText",
               "ClassificationContentMarkingHeaderFontProps",
               "ClassificationContentMarkingHeaderShapeIds",
               "MSIP_Label"]

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def texto_visible(ruta):
    with zipfile.ZipFile(ruta) as z:
        raiz = etree.fromstring(z.read("word/document.xml"))
    return "".join(t.text or "" for t in raiz.iter(f"{W}t"))


print(f"origen : {os.path.basename(ORIGEN)}")
print(f"destino: {os.path.basename(DESTINO)}")
antes = texto_visible(ORIGEN)

with zipfile.ZipFile(ORIGEN) as z:
    partes = z.infolist()
    datos = {i.filename: z.read(i.filename) for i in partes}

# ---------------------------------------------------------------------------
# 1. Diagnostico previo: que hay y donde
# ---------------------------------------------------------------------------
print("\n=== LO QUE SE ENCUENTRA ===")
hay_etiqueta = PARTE_ETIQUETA in datos
print(f"  {PARTE_ETIQUETA}: {'presente' if hay_etiqueta else 'ausente'}")
if hay_etiqueta:
    txt = datos[PARTE_ETIQUETA].decode("utf-8", "ignore")
    bits = re.search(r'contentBits="(\d+)"', txt)
    print(f"    contentBits={bits.group(1) if bits else '-'} "
          "(2 = marcar el contenido en el pie)")

custom = datos.get("docProps/custom.xml", b"").decode("utf-8", "ignore")
presentes = [p for p in PROPIEDADES if f'name="{p}"' in custom]
for p in presentes:
    m = re.search(rf'name="{re.escape(p)}"><vt:lpwstr>([^<]*)</vt:lpwstr>', custom)
    print(f"  propiedad {p} = {m.group(1)[:60] if m else '?'!r}")

# Formas de marcado que pudieran quedar en encabezados y pies
formas = 0
for k, b in datos.items():
    if re.search(r"(header|footer)\d*\.xml$", k):
        t = b.decode("utf-8", "ignore")
        formas += len(re.findall(r"<v:shape\b", t))
        formas += t.count("Internal Use")
print(f"  formas o textos de clasificacion en encabezados y pies: {formas}")

if not hay_etiqueta and not presentes and not formas:
    print("\nno hay nada que retirar: el documento ya esta limpio")
    raise SystemExit(0)

# ---------------------------------------------------------------------------
# 2. Retirada
# ---------------------------------------------------------------------------
print("\n=== RETIRADA ===")
quitar_partes = set()
if hay_etiqueta:
    quitar_partes.add(PARTE_ETIQUETA)
    print(f"  se elimina la parte {PARTE_ETIQUETA}")

# La declaracion de la parte en [Content_Types].xml
ct = datos["[Content_Types].xml"].decode("utf-8")
ct_nuevo = re.sub(r'<Override PartName="/' + re.escape(PARTE_ETIQUETA)
                  + r'"[^>]*/>', "", ct)
if ct_nuevo != ct:
    datos["[Content_Types].xml"] = ct_nuevo.encode("utf-8")
    print("  se retira su declaracion de [Content_Types].xml")

# Y su relacion, si la declara el paquete
for rel in ("_rels/.rels", "word/_rels/document.xml.rels"):
    if rel not in datos:
        continue
    t = datos[rel].decode("utf-8")
    t2 = re.sub(r'<Relationship[^>]*Target="[^"]*LabelInfo\.xml"[^>]*/>', "", t)
    if t2 != t:
        datos[rel] = t2.encode("utf-8")
        print(f"  se retira la relacion declarada en {rel}")

# Las propiedades personalizadas del marcado
if presentes:
    nuevo = custom
    for p in presentes:
        nuevo = re.sub(r"<property\b[^>]*name=\"" + re.escape(p)
                       + r"\"[^>]*>.*?</property>", "", nuevo, flags=re.S)
    # Los pid de las propiedades restantes deben ser consecutivos desde 2
    trozos = re.findall(r"<property\b.*?</property>", nuevo, re.S)
    for k, trozo in enumerate(trozos, start=2):
        renumerado = re.sub(r'pid="\d+"', f'pid="{k}"', trozo)
        nuevo = nuevo.replace(trozo, renumerado)
    datos["docProps/custom.xml"] = nuevo.encode("utf-8")
    print(f"  se eliminan {len(presentes)} propiedades de marcado y se "
          f"renumeran los {len(trozos)} pid restantes")

# ---------------------------------------------------------------------------
# 3. Reescritura del paquete
# ---------------------------------------------------------------------------
temporal = DESTINO + ".tmp"
with zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED) as destino:
    for info in partes:
        if info.filename in quitar_partes:
            continue
        destino.writestr(info, datos[info.filename])
os.replace(temporal, DESTINO)

# ---------------------------------------------------------------------------
# 4. Comprobaciones
# ---------------------------------------------------------------------------
print("\n=== COMPROBACION ===")
with zipfile.ZipFile(DESTINO) as z:
    if z.testzip() is not None:
        raise SystemExit("el paquete quedo corrupto")
    nuevas = z.namelist()
    for k in nuevas:
        if k.endswith(".xml") or k.endswith(".rels"):
            etree.fromstring(z.read(k))
print(f"  paquete valido, {len(nuevas)} partes "
      f"({len(partes) - len(nuevas)} menos que el original)")

with zipfile.ZipFile(DESTINO) as z:
    restos = []
    for k in nuevas:
        if not (k.endswith(".xml") or k.endswith(".rels")):
            continue
        t = z.read(k).decode("utf-8", "ignore")
        for clave in ("Internal Use", "ClassificationContentMarking",
                      "mipLabelMetadata", "LabelInfo"):
            if clave in t:
                restos.append(f"{k}: {clave}")
print(f"  restos de la etiqueta: {restos if restos else 'ninguno'}")
if restos:
    raise SystemExit(1)

despues = texto_visible(DESTINO)
if antes != despues:
    print(f"  CAMBIO EL TEXTO: {len(antes)} -> {len(despues)}")
    raise SystemExit(1)
print(f"  el texto visible es identico ({len(antes)} caracteres)")

a, b = docx.Document(ORIGEN), docx.Document(DESTINO)
with zipfile.ZipFile(ORIGEN) as za, zipfile.ZipFile(DESTINO) as zb:
    ma = {k: za.read(k) for k in za.namelist() if k.startswith("word/media/")}
    mb = {k: zb.read(k) for k in zb.namelist() if k.startswith("word/media/")}
    ca = za.read("word/comments.xml") if "word/comments.xml" in za.namelist() else b""
    cb = zb.read("word/comments.xml") if "word/comments.xml" in zb.namelist() else b""
print(f"  parrafos {len(a.paragraphs)} -> {len(b.paragraphs)}   "
      f"tablas {len(a.tables)} -> {len(b.tables)}   "
      f"imagenes {len(ma)} -> {len(mb)} (identicas: {ma == mb})   "
      f"comentarios intactos: {ca == cb}")
if (len(a.paragraphs), len(a.tables)) != (len(b.paragraphs), len(b.tables)) \
        or ma != mb or ca != cb:
    raise SystemExit("algo mas ha cambiado")

print(f"\nGuardado: {os.path.basename(DESTINO)}")
print("\nAVISO: si tu Word esta gestionado por la organizacion con una politica de")
print("etiquetado obligatorio, puede volver a aplicar la etiqueta al guardar. La")
print("via definitiva es cambiarla desde Word: pestana Inicio, boton Sensibilidad,")
print("y elegir la etiqueta publica o sin clasificar.")
